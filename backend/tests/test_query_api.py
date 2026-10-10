"""/api/query endpoints on a SQLite shop database."""

import time

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine

from app.api import create_app
from app.appdb import AppDB
from app.config import Settings, load_connections
from app.security import hash_password

from .helpers import auth
from .query_data import load

# A query that runs until it is stopped (SQLite counts to a billion).
SLOW_SQL = (
    "WITH RECURSIVE n(i) AS (SELECT 1 UNION ALL SELECT i + 1 FROM n WHERE i < 1000000000) "
    "SELECT count(*) FROM n"
)


def make_client(tmp_path, **settings):
    shop = tmp_path / "shop.db"
    engine = create_engine(f"sqlite:///{shop}")
    load(engine)
    engine.dispose()
    connections = load_connections({"NCODE_CONN_SHOP_ENGINE": "sqlite", "NCODE_CONN_SHOP_DATABASE": str(shop)})
    app_db = AppDB(tmp_path / "app.sqlite")
    app_db.create_user("admin", hash_password("admin-pass-1"), "admin")
    app_db.create_user("alice", hash_password("alice-pass-1"))
    app_db.create_user("bob", hash_password("bob-pass-123"), allowed=("other",))
    defaults = {"search_workers": 2, "query_workers": 4}
    client = TestClient(create_app(Settings(**{**defaults, **settings}), connections, app_db))
    return client, app_db


@pytest.fixture()
def env(tmp_path):
    client, app_db = make_client(tmp_path)
    with client:
        yield client, app_db, auth(client, "alice", "alice-pass-1"), auth(client, "admin", "admin-pass-1")


REVENUE = {
    "connection": "shop",
    "source": {"table": "q_orders"},
    "joins": [{"table": "q_customers"}],
    "aggregations": [{"fn": "sum", "ref": "q_orders.total", "alias": "revenue"}],
    "breakouts": [{"ref": "q_customers.region"}],
    "filters": {"rules": [{"ref": "q_orders.status", "op": "=", "value": "paid"}]},
    "order": [{"ref": "revenue", "dir": "desc"}],
}


def post(client, headers, path, body, status=200):
    res = client.post(path, json=body, headers=headers)
    assert res.status_code == status, res.text
    return res.json()


def wait_done(client, headers, task_id, timeout=10):
    deadline = time.time() + timeout
    while time.time() < deadline:
        data = client.get(f"/api/query/tasks/{task_id}", params={"wait_ms": 500}, headers=headers).json()
        if data["status"] != "running":
            return data
    raise AssertionError("query did not finish")


# ---------------------------------------------------------------- builder questions

def test_run_question(env):
    client, _, alice, _ = env
    data = post(client, alice, "/api/query", {"spec": REVENUE})
    assert data["status"] == "completed", data
    result = data["result"]
    assert [(c["name"], c["type"], c["role"]) for c in result["columns"]] == [
        ("region", "string", "dimension"), ("revenue", "number", "measure"),
    ]
    assert result["rows"] == [["US", 200.0], ["EU", 150.0], [None, 10.0]]
    assert result["stats"]["row_count"] == 3 and result["stats"]["truncated"] is False
    assert "LEFT OUTER JOIN q_customers" in result["sql"] and "'paid'" in result["sql"]


def test_results_are_cached_and_refresh_skips_the_cache(env):
    client, _, alice, _ = env
    first = post(client, alice, "/api/query", {"spec": REVENUE})["result"]["stats"]
    second = post(client, alice, "/api/query", {"spec": REVENUE})["result"]["stats"]
    fresh = post(client, alice, "/api/query", {"spec": REVENUE, "refresh": True})["result"]["stats"]
    assert (first["cached"], second["cached"], fresh["cached"]) == (False, True, False)


def test_truncation_flag(env):
    client, _, alice, _ = env
    spec = {"connection": "shop", "source": {"table": "q_orders"}, "limit": 3}
    result = post(client, alice, "/api/query", {"spec": spec})["result"]
    assert result["stats"]["row_count"] == 3 and result["stats"]["truncated"] is True


def test_compile_shows_sql_without_running(env):
    client, _, alice, _ = env
    data = post(client, alice, "/api/query/compile", {"spec": REVENUE})
    assert data["sql"].startswith("SELECT q_customers.region") and data["limit"] == 2000
    assert data["columns"] == [{"name": "region", "type": "string"}, {"name": "revenue", "type": "number"}]


def test_question_errors_are_400(env):
    client, _, alice, _ = env
    bad = {**REVENUE, "aggregations": [{"fn": "sum", "ref": "q_customers.name"}]}
    assert "sum needs a number" in post(client, alice, "/api/query", {"spec": bad}, status=400)["detail"]
    post(client, alice, "/api/query", {"spec": {**REVENUE, "version": 2}}, status=422)


def test_permissions(env):
    client, _, _, _ = env
    bob = auth(client, "bob", "bob-pass-123")
    post(client, bob, "/api/query", {"spec": REVENUE}, status=404)
    post(client, bob, "/api/query/sql", {"connection": "shop", "sql": "SELECT 1"}, status=404)
    post(client, bob, "/api/query/compile", {"spec": REVENUE}, status=404)
    assert client.post("/api/query", json={"spec": REVENUE}).status_code == 401


# ---------------------------------------------------------------- SQL mode

def test_sql_with_params_and_percent(env):
    client, _, alice, _ = env
    body = {
        "connection": "shop",
        "sql": "SELECT name, region FROM q_customers WHERE region = {{region}} AND name NOT LIKE '50%' ORDER BY name;",
        "params": {"region": "EU", "unused": 1},
    }
    data = post(client, alice, "/api/query/sql", body)
    assert data["status"] == "completed", data
    assert data["result"]["rows"] == [["Anna", "EU"]]
    assert data["result"]["sql"] == body["sql"]


def test_sql_limit_and_truncation(env):
    client, _, alice, _ = env
    data = post(client, alice, "/api/query/sql", {"connection": "shop", "sql": "SELECT id FROM q_orders ORDER BY id", "limit": 2})
    assert data["result"]["rows"] == [[1], [2]] and data["result"]["stats"]["truncated"] is True


def test_sql_guard_and_params_errors(env):
    client, _, alice, admin = env
    assert "Only SELECT" in post(client, alice, "/api/query/sql", {"connection": "shop", "sql": "DELETE FROM q_orders"}, status=400)["detail"]
    assert "readfile" in post(client, alice, "/api/query/sql", {"connection": "shop", "sql": "SELECT readfile('/etc/passwd')"}, status=400)["detail"].lower()
    assert "{{r}}" in post(client, alice, "/api/query/sql", {"connection": "shop", "sql": "SELECT {{r}}"}, status=400)["detail"]
    post(client, alice, "/api/query/sql", {"connection": "shop", "sql": "SELECT name FROM sqlite_master"}, status=400)
    assert post(client, admin, "/api/query/sql", {"connection": "shop", "sql": "SELECT name FROM sqlite_master"})["status"] == "completed"


def test_database_errors_are_reported_on_the_task(env):
    client, _, alice, _ = env
    data = post(client, alice, "/api/query/sql", {"connection": "shop", "sql": "SELECT no_such_column FROM q_orders"})
    assert data["status"] == "error" and "no_such_column" in data["error"]


def test_sql_check(env):
    client, _, alice, _ = env
    good = post(client, alice, "/api/query/sql/check", {"connection": "shop", "sql": "SELECT * FROM q_orders WHERE id = {{id}}"})
    assert good == {"ok": True, "error": None, "params": ["id"], "tables": ["q_orders"]}
    badc = post(client, alice, "/api/query/sql/check", {"connection": "shop", "sql": "SELECT 1; SELECT 2"})
    assert badc["ok"] is False and "one statement" in badc["error"]


# ---------------------------------------------------------------- tasks, cancel, limits, audit

def test_cancel_a_running_query(env):
    client, _, alice, admin = env
    data = post(client, alice, "/api/query/sql", {"connection": "shop", "sql": SLOW_SQL, "wait_ms": 0})
    assert data["status"] == "running"
    task_id = data["task_id"]
    other = auth(client, "bob", "bob-pass-123")
    assert client.delete(f"/api/query/tasks/{task_id}", headers=other).status_code == 404  # not bob's
    assert client.delete(f"/api/query/tasks/{task_id}", headers=alice).json()["status"] in ("cancelling", "cancelled")
    assert wait_done(client, alice, task_id)["status"] == "cancelled"
    # the connection is usable afterwards
    assert post(client, alice, "/api/query/sql", {"connection": "shop", "sql": "SELECT 1 AS one"})["result"]["rows"] == [[1]]


def test_timeout_stops_the_query(tmp_path):
    client, _ = make_client(tmp_path, statement_timeout_sec=1)
    with client:
        alice = auth(client, "alice", "alice-pass-1")
        data = post(client, alice, "/api/query/sql", {"connection": "shop", "sql": SLOW_SQL, "wait_ms": 0})
        done = wait_done(client, alice, data["task_id"], timeout=15)
        assert done["status"] == "error" and "longer than 1 seconds" in done["error"]


def test_per_user_concurrency_limit(tmp_path):
    client, _ = make_client(tmp_path, max_concurrent_queries=1)
    with client:
        alice = auth(client, "alice", "alice-pass-1")
        first = post(client, alice, "/api/query/sql", {"connection": "shop", "sql": SLOW_SQL, "wait_ms": 0})
        assert "already have 1" in post(client, alice, "/api/query/sql", {"connection": "shop", "sql": "SELECT 1"}, status=429)["detail"]
        client.delete(f"/api/query/tasks/{first['task_id']}", headers=alice)
        wait_done(client, alice, first["task_id"])
        assert post(client, alice, "/api/query/sql", {"connection": "shop", "sql": "SELECT 1"})["status"] == "completed"


def test_runs_are_audited(env):
    client, app_db, alice, _ = env
    post(client, alice, "/api/query", {"spec": REVENUE})
    post(client, alice, "/api/query/sql", {"connection": "shop", "sql": "SELECT 1"})
    time.sleep(0.2)
    entries = {e.action: e for e in app_db.list_audit(10)}
    assert entries["query"].status == "completed" and entries["query"].found == 3
    assert entries["sql"].phrase == "SELECT 1" and entries["sql"].connection == "shop"


def test_limits_endpoint_reports_query_caps(env):
    client, _, alice, _ = env
    limits = client.get("/api/limits", headers=alice).json()
    assert limits["query_default_rows"] == 2000 and limits["query_max_rows"] == 10000


# ---------------------------------------------------------------- Excel export

def test_xlsx_export_keeps_types_and_neutralises_formulas(env):
    import io

    import openpyxl  # test-only reader

    client, app_db, alice, _ = env
    sql = (
        "SELECT id, total, created_at, paid_on, status, '=1+1' AS tricky FROM q_orders "
        "WHERE id IN (1, 6) ORDER BY id"
    )
    data = post(client, alice, "/api/query/sql", {"connection": "shop", "sql": sql})
    res = client.get(f"/api/query/tasks/{data['task_id']}/export.xlsx", headers=alice)
    assert res.status_code == 200 and res.headers["content-type"].startswith("application/vnd.openxmlformats")
    sheet = openpyxl.load_workbook(io.BytesIO(res.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("id", "total", "created_at", "paid_on", "status", "tricky")
    assert rows[1][0] == 1 and rows[1][1] == 100
    assert rows[2][4] is None  # NULL stays an empty cell
    assert rows[1][5] == "=1+1" and sheet["F2"].data_type == "s"  # text, not a formula
    assert any(e.action == "export_xlsx" for e in app_db.list_audit(5))


def test_xlsx_export_needs_a_finished_task_of_your_own(env):
    client, _, alice, admin = env
    running = post(client, alice, "/api/query/sql", {"connection": "shop", "sql": SLOW_SQL, "wait_ms": 0})
    assert client.get(f"/api/query/tasks/{running['task_id']}/export.xlsx", headers=alice).status_code == 409
    client.delete(f"/api/query/tasks/{running['task_id']}", headers=alice)
    other = auth(client, "bob", "bob-pass-123")
    assert client.get(f"/api/query/tasks/{running['task_id']}/export.xlsx", headers=other).status_code == 404


def test_xlsx_dates_from_a_question_are_real_dates(env):
    import datetime as dt
    import io

    import openpyxl

    client, _, alice, _ = env
    spec = {"connection": "shop", "source": {"table": "q_orders"}, "aggregations": [{"fn": "count", "alias": "n"}],
            "breakouts": [{"ref": "created_at", "bucket": "month", "alias": "month"}]}
    data = post(client, alice, "/api/query", {"spec": spec})
    content = client.get(f"/api/query/tasks/{data['task_id']}/export.xlsx", headers=alice).content
    first = list(openpyxl.load_workbook(io.BytesIO(content)).active.iter_rows(min_row=2, values_only=True))[0]
    assert first == (dt.datetime(2025, 12, 1), 1)
