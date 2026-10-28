"""Uploaded files ("datasets", roadmap M6): CSV, Excel and Parquet.

An upload is read once, its column types are inferred, and it is kept as Ncode's own
Parquet copy in data/datasets/. The original file is not kept. Datasets live in
collections, like questions, so the same view / edit permissions apply; they are
used as sources of multi-source questions.
"""

from __future__ import annotations

import codecs
import datetime as dt
import json
import logging
import os
import re
import time
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

from sqlalchemy import delete, func, insert, or_, select, update

from ..appdb import User, users_table
from ..content.store import ContentError, ContentStore, Level, NotFound, Snapshot, _clean_name
from ..content.tables import datasets_table
from .workspace import Workspace, literal, ncode_type, quote, unique_names

logger = logging.getLogger(__name__)

Format = Literal["csv", "xlsx", "parquet"]
FILE_RE = re.compile(r"^ds_[0-9a-f]{32}\.parquet$")
TEMP_RE = re.compile(r"^upload_[0-9a-f]{32}[.\w]*$")  # uploads being read (app/federation/routes.py) and their UTF-8 copies
EXTENSIONS: dict[str, Format] = {
    ".csv": "csv", ".tsv": "csv", ".txt": "csv", ".xlsx": "xlsx", ".xlsm": "xlsx", ".parquet": "parquet",
}
DELIMITERS = {"auto": None, "comma": ",", "semicolon": ";", "tab": "\t", "pipe": "|"}
# "auto": UTF-16 / UTF-8 by their marks, UTF-8 when the whole file is valid UTF-8, else Windows-1251 when
# the text looks Cyrillic (Russian Excel exports), else Windows-1252 (Western European Excel exports)
ENCODINGS = ("auto", "utf-8", "windows-1251", "windows-1252", "latin-1", "utf-16")
_DUCKDB_ENCODINGS = {"utf-8", "latin-1", "utf-16"}  # the others are copied to UTF-8 first
_CODECS = {"utf-8": "utf-8-sig", "windows-1251": "cp1251", "windows-1252": "cp1252", "latin-1": "latin-1", "utf-16": "utf-16"}
# "comma": 1 234,5 (decimal comma, spaces or non-breaking spaces between thousands); "auto" picks it
# for ;-separated files whose numbers use commas, like Russian and European Excel exports
NUMBER_FORMATS = ("auto", "point", "comma")
_SAMPLE_BYTES = 256 * 1024
_COMMA_NUMBER = re.compile(r"[-+]?\d{1,3}(?:[ \u00a0\u202f]\d{3})+(?:,\d+)?|[-+]?\d+,\d+")
_POINT_NUMBER = re.compile(r"[-+]?\d+\.\d+")
_NBSP_IN_NUMBER = re.compile(r"(?<=\d)[\u00a0\u202f](?=\d)")
_SIMPLE_TYPES = ("number", "string", "boolean", "date", "datetime", "time")


class IngestError(ValueError):
    """The file cannot be read as asked; shown to the user."""


@dataclass
class IngestOptions:
    delimiter: str = "auto"          # key of DELIMITERS
    header: bool = True              # first row holds the column names
    encoding: str = "auto"           # one of ENCODINGS
    sheet: str | None = None         # Excel: sheet name (default: the first)
    number_format: str = "auto"      # one of NUMBER_FORMATS


@dataclass
class Ingested:
    file: str
    columns: list[dict[str, str]]
    row_count: int
    size_bytes: int
    read_as: dict[str, str] = field(default_factory=dict)  # CSV: the encoding and number format used


def format_for(filename: str) -> Format:
    ext = Path(filename).suffix.lower()
    if ext == ".xls":
        raise IngestError("Old .xls workbooks cannot be read: save the file as .xlsx or .csv")
    fmt = EXTENSIONS.get(ext)
    if fmt is None:
        raise IngestError("Upload a .csv, .tsv, .xlsx or .parquet file")
    return fmt


def ingest(
    src: Path, fmt: Format, options: IngestOptions, *, dest_dir: Path, max_rows: int, memory_mb: int, threads: int
) -> Ingested:
    """Read an uploaded file into a new Parquet file in dest_dir."""
    if options.delimiter not in DELIMITERS:
        raise IngestError(f"Unknown delimiter {options.delimiter!r}")
    if options.encoding not in ENCODINGS:
        raise IngestError(f"Unknown encoding {options.encoding!r}")
    if options.number_format not in NUMBER_FORMATS:
        raise IngestError(f"Unknown number format {options.number_format!r}")
    if options.number_format == "comma" and options.delimiter == "comma":
        raise IngestError("A comma cannot be both the separator and the decimal mark: choose another separator")
    dest_dir.mkdir(parents=True, exist_ok=True)
    name = f"ds_{uuid.uuid4().hex}.parquet"
    target = dest_dir / name
    copy: Path | None = None
    read_as: dict[str, str] = {}
    done = False
    # This session reads and writes files, so it never runs SQL written by a user;
    # every value in it is one of the checked options above or a path Ncode made.
    ws = Workspace(memory_mb=memory_mb, threads=threads)
    try:
        if fmt == "csv":
            encoding = detect_encoding(src) if options.encoding == "auto" else options.encoding
            number = options.number_format
            if number == "auto":
                number = sniff_number_format(_sample_text(src, encoding), DELIMITERS[options.delimiter])
            read_as = {"encoding": encoding, "number_format": number}
            path = src
            if encoding not in _DUCKDB_ENCODINGS or number == "comma":
                copy = path = dest_dir / f"upload_{uuid.uuid4().hex}.utf8.csv"
                _to_utf8(src, encoding, copy, number == "comma")
                encoding = "utf-8"
            args = [literal(str(path)), "auto_detect = true", "sample_size = -1", f"header = {'true' if options.header else 'false'}",
                    f"encoding = {literal(encoding)}", "null_padding = true"]
            delim = DELIMITERS[options.delimiter]
            if delim is not None:
                args.append(f"delim = {literal(delim)}")
            if number == "comma":
                args += ["decimal_separator = ','", "thousands = ' '"]
            try:
                _create(ws, f"SELECT * FROM read_csv({', '.join(args)}) LIMIT {int(max_rows) + 1}")
            except Exception as exc:
                # DuckDB's parallel reader cannot pad short rows (null_padding) when a quoted value holds a
                # line break; the single-threaded reader can
                if "parallel" not in str(exc).lower():
                    raise
                args.append("parallel = false")
                _create(ws, f"SELECT * FROM read_csv({', '.join(args)}) LIMIT {int(max_rows) + 1}")
        elif fmt == "parquet":
            _create(ws, f"SELECT * FROM read_parquet({literal(str(src))}) LIMIT {int(max_rows) + 1}")
        else:
            names, rows = _read_xlsx(src, options, max_rows)
            ws.load_rows("data", names, rows)
            _dates_without_time(ws)
        rows = ws.count("data")
        if rows > max_rows:
            raise IngestError(f"The file has more than {max_rows:,} rows (the limit for uploads)")
        if rows == 0:
            raise IngestError("The file has no data rows")
        _plain_types(ws)
        ws.execute(f"COPY data TO {literal(str(target))} (FORMAT parquet, COMPRESSION zstd)")
        columns = [{"name": n, "type": ncode_type(t)} for n, t in ws.columns("data")]
        done = True
    except IngestError:
        raise
    except UnicodeDecodeError:
        raise IngestError(_NOT_UTF8 if read_as.get("encoding") == "utf-8" else
                          f"The file is not valid {read_as.get('encoding', 'text')}: choose another encoding") from None
    except Exception as exc:
        raise IngestError(friendly_error(exc, fmt, [src, target, dest_dir] + ([copy] if copy else []))) from None
    finally:
        ws.close()  # DuckDB lets go of the files before they are removed (Windows cannot remove open files)
        if copy is not None:
            remove_path(copy)
        if not done:
            remove_path(target)
    return Ingested(file=name, columns=columns, row_count=rows, size_bytes=target.stat().st_size, read_as=read_as)


# ---------------------------------------------------------------- text files: encodings and numbers

def detect_encoding(src: Path) -> str:
    with open(src, "rb") as fh:
        head = fh.read(_SAMPLE_BYTES)
    if head.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    if _is_utf8(src):
        return "utf-8"
    return "windows-1251" if looks_cyrillic(head.decode("cp1251", errors="replace")) else "windows-1252"


def _is_utf8(src: Path) -> bool:
    decoder = codecs.getincrementaldecoder("utf-8")()
    try:
        with open(src, "rb") as fh:
            while chunk := fh.read(1 << 20):
                decoder.decode(chunk)
        decoder.decode(b"", final=True)
    except UnicodeDecodeError:
        return False
    return True


def looks_cyrillic(text: str) -> bool:
    """Most words with non-ASCII letters are all Cyrillic (Western text read as Windows-1251 gives
    words that mix Latin and Cyrillic letters, like "cafй")."""
    words = [w for w in re.findall(r"[^\W\d_]+", text) if not w.isascii()]
    if not words:
        return False
    cyrillic = sum(1 for w in words if all("\u0400" <= ch <= "\u04ff" for ch in w))
    return cyrillic >= 0.6 * len(words)


def _sample_text(src: Path, encoding: str) -> str:
    with open(src, "rb") as fh:
        raw = fh.read(_SAMPLE_BYTES)
    text = codecs.getincrementaldecoder(_CODECS[encoding])(errors="replace").decode(raw)
    return text if len(raw) < _SAMPLE_BYTES else text[: text.rfind("\n") + 1]


def sniff_number_format(text: str, delimiter: str | None) -> str:
    """"comma" for ;-separated text whose numbers use a decimal comma and none a decimal point."""
    lines = [line for line in text.splitlines() if line.strip()][:500]
    if not lines:
        return "point"
    if delimiter is None:
        first = lines[0]
        delimiter = max((",", ";", "\t", "|"), key=first.count)
    if delimiter != ";":
        return "point"
    comma = point = 0
    for line in lines:
        for cell in line.split(";"):
            value = cell.strip().strip('"').strip()
            if _COMMA_NUMBER.fullmatch(value):
                comma += 1
            elif _POINT_NUMBER.fullmatch(value):
                point += 1
    return "comma" if comma and not point else "point"


def _to_utf8(src: Path, encoding: str, dest: Path, number_spaces: bool) -> None:
    """A UTF-8 copy of a text file, line by line (memory stays flat; the copy is at most 3x the upload)."""
    errors = "strict" if encoding in ("utf-8", "utf-16") else "replace"  # Windows-1251 has one unused byte
    with open(src, encoding=_CODECS[encoding], errors=errors, newline="") as fin, \
            open(dest, "w", encoding="utf-8", newline="") as fout:
        for line in fin:
            fout.write(_NBSP_IN_NUMBER.sub(" ", line) if number_spaces else line)


# ---------------------------------------------------------------- errors shown to the user

_NOT_UTF8 = ("The file is not UTF-8 text: choose its encoding, for example Windows-1251 (Cyrillic) "
             "or Windows-1252 (Western European)")
_QUOTED_PATH = re.compile(r"""(["'])(?:[A-Za-z]:)?[\\/][^"'\n]*\1""")


def friendly_error(exc: BaseException, fmt: str, paths: list[Path]) -> str:
    """A sentence for the user from a DuckDB / openpyxl error: no server paths, no reader internals."""
    text = str(exc).strip()
    for p in sorted({str(p) for p in paths}, key=len, reverse=True):
        text = text.replace(p, "the file")
    text = _QUOTED_PATH.sub("the file", text)
    line = re.search(r"CSV Error on Line: (\d+)", text)
    where = f"line {line.group(1)}" if line else "a line"
    counts = re.search(r"Expected Number of Columns: (\d+) Found: (\d+)", text)
    if counts:
        return (f"Could not read the file as CSV: {where} has {counts.group(2)} values, the header has "
                f"{counts.group(1)}. Check the separator.")
    low = text.lower()
    if "invalid unicode" in low or "not utf-8 encoded" in low:
        return _NOT_UTF8
    if "unterminated quote" in low or ("quote" in low and "csv error" in low):
        return f'Could not read the file as CSV: a quoted value on {where} is not closed (a " is missing).'
    if fmt == "parquet" and ("parquet" in low or "magic" in low):
        return "The file is not a valid Parquet file"
    if "out of memory" in low or "memory limit" in low:
        return "The file needs more memory than this server allows for uploads: split it or ask an administrator"
    first = next((ln for ln in text.splitlines() if ln.strip()), type(exc).__name__)
    first = re.sub(r"^[A-Za-z ]*(Error|Exception):\s*", "", first.strip())
    return f"Could not read the file: {first}"[:500]


def _create(ws: Workspace, select_sql: str) -> None:
    ws.execute(f"CREATE TABLE data AS {select_sql}")
    # header names exactly as DuckDB read them, made unique case-insensitively
    names = [n for n, _ in ws.columns("data")]
    for old, new in zip(names, unique_names(names)):
        if old != new:
            ws.execute(f"ALTER TABLE data RENAME COLUMN {quote(old)} TO {quote(new)}")


def _plain_types(ws: Workspace) -> None:
    """Lists, structs, blobs ... become text: the builder and the charts work with simple types."""
    for name, type_ in ws.columns("data"):
        if ncode_type(type_) not in _SIMPLE_TYPES:
            ws.execute(f"ALTER TABLE data ALTER {quote(name)} TYPE VARCHAR USING CAST({quote(name)} AS VARCHAR)")


def _dates_without_time(ws: Workspace) -> None:
    """Excel has no date type: date-times that are all at midnight are dates."""
    for name, type_ in ws.columns("data"):
        if type_ == "TIMESTAMP":
            col = quote(name)
            only_dates = ws.execute(
                f"SELECT bool_and(CAST({col} AS TIME) = TIME '00:00:00') FROM data WHERE {col} IS NOT NULL"
            ).fetchone()[0]
            if only_dates:
                ws.execute(f"ALTER TABLE data ALTER {col} TYPE DATE USING CAST({col} AS DATE)")


def _read_xlsx(src: Path, options: IngestOptions, max_rows: int) -> tuple[list[str], list[tuple]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(src, read_only=True, data_only=True)
    except Exception as exc:
        raise IngestError(f"Could not open the workbook: {str(exc).splitlines()[0] if str(exc) else type(exc).__name__}") from None
    try:
        if options.sheet:
            if options.sheet not in wb.sheetnames:
                raise IngestError(f"The workbook has no sheet {options.sheet!r} (sheets: {', '.join(wb.sheetnames)})")
            sheet = wb[options.sheet]
        else:
            sheet = wb.worksheets[0]
        header: list[Any] | None = None
        rows: list[tuple] = []
        width = 0
        for raw in sheet.iter_rows(values_only=True):
            if raw is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in raw):
                continue  # empty rows
            values = [v.strip() if isinstance(v, str) else v for v in raw]
            if header is None:
                if options.header:
                    header = values
                    width = max((i + 1 for i, v in enumerate(values) if v not in (None, "")), default=0)
                    continue
                header = []
            width = max(width, max((i + 1 for i, v in enumerate(values) if v not in (None, "")), default=0))
            rows.append(tuple(values))
            if len(rows) > max_rows:
                raise IngestError(f"The sheet has more than {max_rows:,} rows (the limit for uploads)")
    finally:
        wb.close()
    if header is None or width == 0:
        raise IngestError("The sheet is empty")
    names = [str(header[i]) if i < len(header) and header[i] not in (None, "") else f"column_{i + 1}" for i in range(width)]
    fixed = [tuple(_cell(r[i]) if i < len(r) else None for i in range(width)) for r in rows]
    return names, fixed


def _cell(v: Any) -> Any:
    if isinstance(v, str) and v == "":
        return None
    if isinstance(v, dt.timedelta):  # durations: seconds
        return v.total_seconds()
    return v


# ---------------------------------------------------------------- the store

@dataclass
class Dataset:
    id: int
    name: str
    description: str | None
    collection_id: int
    owner: str | None
    archived: bool
    file: str
    columns: list[dict[str, str]]
    row_count: int
    size_bytes: int
    source_name: str | None
    source_format: str
    created_at: float
    updated_at: float
    updated_by: str | None
    access: Level | None = None
    extra: dict[str, Any] = field(default_factory=dict)

    def summary(self) -> dict[str, Any]:
        return {
            "id": self.id, "name": self.name, "description": self.description, "collection_id": self.collection_id,
            "collection": self.extra.get("collection"), "owner": self.owner, "archived": self.archived,
            "columns": self.columns, "row_count": self.row_count, "size_bytes": self.size_bytes,
            "source_name": self.source_name, "source_format": self.source_format, "access": self.access,
            "created_at": self.created_at, "updated_at": self.updated_at, "updated_by": self.updated_by,
        }


class DatasetStore:
    def __init__(self, content: ContentStore, directory: Path):
        self.content = content
        self.engine = content.engine
        self.directory = directory

    def path(self, d: Dataset) -> Path:
        if not FILE_RE.match(d.file):  # never follow a name Ncode did not make
            raise ContentError("The dataset's file name is not valid", 500)
        return self.directory / d.file

    def _select(self):
        return select(datasets_table, users_table.c.username.label("owner")).outerjoin(
            users_table, users_table.c.id == datasets_table.c.owner_id
        )

    def _to(self, row: Any, snap: Snapshot, user: User) -> Dataset:
        col = snap.collections.get(row.collection_id)
        try:
            columns = json.loads(row.columns)
        except ValueError:
            columns = []
        return Dataset(
            id=row.id, name=row.name, description=row.description, collection_id=row.collection_id, owner=row.owner,
            archived=bool(row.archived), file=row.file, columns=columns, row_count=row.row_count,
            size_bytes=row.size_bytes, source_name=row.source_name, source_format=row.source_format,
            created_at=row.created_at, updated_at=row.updated_at, updated_by=row.updated_by,
            access=snap.access(user, row.collection_id),
            extra={"collection": {"id": col.id, "name": col.name} if col else None},
        )

    def get(self, user: User, did: int, level: Level = "view") -> Dataset:
        snap = self.content.snapshot()
        with self.engine.connect() as conn:
            row = conn.execute(self._select().where(datasets_table.c.id == did)).first()
        if row is None or snap.access(user, row.collection_id) is None:
            raise NotFound("Dataset")
        self.content._require(snap, user, row.collection_id, level)
        d = self._to(row, snap, user)
        d.extra["path"] = self.content.breadcrumbs(user, snap, d.collection_id)
        return d

    def list(
        self, user: User, *, collection_id: int | None = None, text: str | None = None, archived: bool = False,
        limit: int = 200,
    ) -> list[Dataset]:
        snap = self.content.snapshot()
        stmt = self._select()
        if collection_id is not None:
            self.content._require(snap, user, collection_id, "view")
            stmt = stmt.where(datasets_table.c.collection_id == collection_id)
        if text:
            pattern = f"%{text.strip().lower()}%"
            stmt = stmt.where(or_(
                func.lower(datasets_table.c.name).like(pattern),
                func.lower(func.coalesce(datasets_table.c.description, "")).like(pattern),
                func.lower(func.coalesce(datasets_table.c.source_name, "")).like(pattern),
            ))
        stmt = stmt.where(datasets_table.c.archived.is_(archived)).order_by(datasets_table.c.updated_at.desc())
        with self.engine.connect() as conn:
            rows = conn.execute(stmt).all()
        out = []
        for row in rows:
            if snap.access(user, row.collection_id) is None:
                continue
            if not archived and collection_id is None and snap.hidden(row.collection_id):
                continue
            out.append(self._to(row, snap, user))
            if len(out) >= limit:
                break
        return out

    def check_target(self, user: User, collection_id: int) -> None:
        self.content._require(self.content.snapshot(), user, collection_id, "edit")

    def create(
        self, user: User, *, name: str, description: str | None, collection_id: int, ingested: Ingested,
        source_name: str | None, source_format: str,
    ) -> Dataset:
        self.check_target(user, collection_id)
        now = time.time()
        with self.engine.begin() as conn:
            did = int(conn.execute(insert(datasets_table).values(
                name=_clean_name(name, 200), description=(description or "").strip() or None,
                collection_id=collection_id, owner_id=user.id, archived=False, file=ingested.file,
                columns=json.dumps(ingested.columns), row_count=ingested.row_count, size_bytes=ingested.size_bytes,
                source_name=(source_name or "")[:255] or None, source_format=source_format,
                created_at=now, updated_at=now, updated_by=user.username,
            )).inserted_primary_key[0])
        return self.get(user, did)

    def update(self, user: User, did: int, **changes: Any) -> Dataset:
        current = self.get(user, did, "edit")
        values: dict[str, Any] = {}
        if "name" in changes and changes["name"] is not None:
            values["name"] = _clean_name(changes["name"], 200)
        if "description" in changes:
            values["description"] = (changes["description"] or "").strip() or None
        if changes.get("collection_id") is not None and changes["collection_id"] != current.collection_id:
            self.check_target(user, changes["collection_id"])
            values["collection_id"] = changes["collection_id"]
        if changes.get("archived") is not None:
            values["archived"] = bool(changes["archived"])
        if values:
            values.update(updated_at=time.time(), updated_by=user.username)
            with self.engine.begin() as conn:
                conn.execute(update(datasets_table).where(datasets_table.c.id == did).values(**values))
        return self.get(user, did)

    def delete(self, user: User, did: int) -> None:
        d = self.get(user, did, "edit")
        if not d.archived:
            raise ContentError("Archive the dataset first; only archived datasets can be deleted", 409)
        with self.engine.begin() as conn:
            conn.execute(delete(datasets_table).where(datasets_table.c.id == did))
        remove_files(self.directory, [d.file])


def remove_files(directory: Path, files: list[str]) -> None:
    for name in files:
        if FILE_RE.match(name):
            remove_path(directory / name)


def remove_path(path: Path, *, attempts: int = 5) -> bool:
    """Remove a file; False when it is still there. On Windows a file another process (or a DuckDB
    session still reading it) has open cannot be removed, so this tries a few times; what is left
    over is removed by sweep_orphans when Ncode starts."""
    for attempt in range(attempts):
        try:
            os.remove(path)
            return True
        except FileNotFoundError:
            return True
        except OSError:
            if attempt + 1 < attempts:
                time.sleep(0.1 * (attempt + 1))
    logger.warning("Could not remove %s; it is removed when Ncode starts next", path.name)
    return False


def sweep_orphans(engine: Any, directory: Path, *, min_age_s: float = 3600) -> int:
    """Remove dataset files no dataset uses and left-over upload copies (older than min_age_s, so a
    file another worker is uploading right now stays). Returns how many were removed."""
    if not directory.is_dir():
        return 0
    with engine.connect() as conn:
        used = {r[0] for r in conn.execute(select(datasets_table.c.file))}
    removed = 0
    now = time.time()
    for path in directory.iterdir():
        name = path.name
        if not (FILE_RE.match(name) and name not in used) and not TEMP_RE.match(name):
            continue
        try:
            if now - path.stat().st_mtime < min_age_s:
                continue
        except OSError:
            continue
        removed += remove_path(path, attempts=1)
    return removed
